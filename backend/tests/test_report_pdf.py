import io
import pytest
from app.api.v1.routes import reports as rp
from app.core import report_export as rx

ROWS = [{"date": "2026-10-01", "orders": 3, "revenue": 1234.5, "note": "বাংলা"}, {"date": "2026-10-02", "orders": 0, "revenue": 99.0}]
CO = {"name": "ABO Enterprise", "address": "Sylhet", "phone": "01700000000", "email": "a@b.c", "site": "https://www.aboenterprise.com"}


@pytest.mark.parametrize("kind", ["sales", "customers", "profit"])
def test_pdf_and_xlsx_build(kind):
    pdf = rx.build_report_pdf(kind, ROWS, "2026-10-01 to 2026-10-02", CO)
    assert pdf[:4] == b"%PDF"
    x = rx.build_report_xlsx(kind, ROWS, "2026-10-01 to 2026-10-02", CO)
    assert x[:2] == b"PK"  # zip container
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x)).active
    assert ws.cell(row=6, column=1).value == "Date" and ws.cell(row=9, column=1).value == "Total"


def test_empty_report_still_builds():
    assert rx.build_report_pdf("sales", [], "p", CO)[:4] == b"%PDF"
    assert rx.build_report_xlsx("sales", [], "p", CO)[:2] == b"PK"


@pytest.mark.asyncio
async def test_routes_return_files(monkeypatch):
    async def fake(db, s, e):
        return ROWS
    monkeypatch.setitem(rp.REPORT_BUILDERS, "sales", fake)
    async def fake_co(db):
        return CO
    monkeypatch.setattr(rx, "company_info", fake_co)
    r = await rp.export_report_pdf("sales", None, None, db=None, _admin="a")
    assert (b"".join([c async for c in r.body_iterator]))[:4] == b"%PDF" and r.headers["content-disposition"].endswith(".pdf")
    r = await rp.export_report_xlsx("sales", None, None, db=None, _admin="a")
    assert r.headers["content-disposition"].endswith(".xlsx")
