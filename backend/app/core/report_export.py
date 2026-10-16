"""Professional report exports (A4 PDF with logo + QR, and styled XLSX).

Reuses what the invoice already ships: the bundled FreeSans fonts (Bangla + ৳ via
HarfBuzz shaping), logo.jpg and the QR helper. Header/footer are drawn on every
page; a summary strip shows the record count and totals of the money columns.
"""
from __future__ import annotations

import io
import re
import secrets
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

BRAND = (29, 78, 216)
BRAND_DARK = (30, 58, 138)
ROW_ALT = (243, 246, 253)
GRID = (219, 227, 240)
MUTED = (107, 114, 128)
INK = (17, 24, 39)

TITLES: dict[str, tuple[str, str]] = {
    "sales": ("Sales Report", "বিক্রয় রিপোর্ট"),
    "revenue": ("Revenue Report", "আয়ের রিপোর্ট"),
    "products": ("Products Report", "পণ্যের রিপোর্ট"),
    "services": ("Services Report", "সেবার রিপোর্ট"),
    "customers": ("Customers Report", "গ্রাহকের রিপোর্ট"),
    "payments": ("Payments Report", "পেমেন্টের রিপোর্ট"),
    "courier": ("Courier Report", "কুরিয়ারের রিপোর্ট"),
    "profit": ("Profit Report", "লাভের রিপোর্ট"),
    "reconciliation": ("Reconciliation Report", "হিসাব মিলানোর রিপোর্ট"),
}
_MONEY = re.compile(r"(revenue|total|amount|profit|cost|price|value|sales|paid|refund|commission|expense)", re.I)
_FONTS = Path(__file__).resolve().parents[2] / "assets" / "fonts"


def _logo() -> Path | None:
    from app.core.invoice import _find_logo_path

    return _find_logo_path()


async def company_info(db: AsyncSession) -> dict[str, str]:
    """Name / address / phone / email — the SAME admin-editable settings the invoice uses
    (InvoiceService._company_info), so changing them in Admin → Settings changes invoices
    and reports together. Adds the site URL (for the QR) and the tagline."""
    from app.core.invoice import _COMPANY_TAGLINE, InvoiceService

    info = dict(await InvoiceService(db)._company_info())
    info["tagline"] = _COMPANY_TAGLINE
    info["site"] = "https://www.aboenterprise.com"
    try:
        from app.models.models import Setting

        row = (await db.execute(select(Setting).where(Setting.key == "site_url", Setting.is_deleted == False))).scalar_one_or_none()  # noqa: E712
        if row and (row.value or "").strip():
            info["site"] = row.value.strip().rstrip("/")
    except Exception:  # noqa: BLE001
        pass
    return info


def _num(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        return float(v)
    return None


def _fmt(key: str, v: Any) -> str:
    if v is None:
        return ""
    n = _num(v)
    if n is None:
        return str(v)
    if _MONEY.search(key):
        return f"৳{n:,.2f}"
    return f"{int(n):,}" if float(n).is_integer() else f"{n:,.2f}"


def _headers(rows: list[dict[str, Any]]) -> list[str]:
    out: list[str] = []
    for r in rows:
        for k in r:
            if k not in out and k != "summary":
                out.append(k)
    return out or ["value"]


def _pretty(key: str) -> str:
    return key.replace("_", " ").title()


def _totals(headers: list[str], rows: list[dict[str, Any]]) -> list[tuple[str, float]]:
    out = []
    for h in headers:
        if _MONEY.search(h) or h in ("orders", "bookings", "units_sold", "order_count", "count"):
            vals = [n for r in rows if (n := _num(r.get(h))) is not None]
            if vals:
                out.append((h, sum(vals)))
    return out[:4]


def report_ref() -> str:
    return f"RPT-{datetime.now(timezone.utc):%Y%m%d}-{secrets.token_hex(2).upper()}"


def build_report_pdf(report_type: str, rows: list[dict[str, Any]], period: str, company: dict[str, str]) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    title_en, title_bn = TITLES.get(report_type, (f"{_pretty(report_type)} Report", "রিপোর্ট"))
    headers = _headers(rows)
    ref = report_ref()
    generated = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    landscape = len(headers) > 6
    logo = _logo()

    from app.core.invoice import _qr_image

    qr = _qr_image(company.get("site", ""))

    class Doc(FPDF):
        def header(self) -> None:  # slim running header from page 2 on; page 1 draws the big band itself
            if self.page_no() == 1:
                return
            self.set_font("rpt", "B", 8)
            self.set_text_color(*BRAND_DARK)
            self.set_xy(M, 20)
            self.cell(0, 10, f"{company.get('name', '')} · {title_en}", new_x="LMARGIN", new_y="NEXT")
            self.set_draw_color(*GRID)
            self.line(M, 34, self.w - M, 34)
            self.set_y(42)

        def footer(self) -> None:
            self.set_y(-46)
            self.set_draw_color(*GRID)
            self.line(M, self.get_y(), self.w - M, self.get_y())
            self.set_font("rpt", "", 7)
            self.set_text_color(*MUTED)
            contact = " · ".join(x for x in (company.get("address"), company.get("phone"), company.get("email")) if x)
            self.set_xy(M, self.get_y() + 4)
            self.multi_cell(self.w - 2 * M - 52, 9, contact or company.get("site", ""), new_x="LMARGIN", new_y="NEXT")
            self.set_x(M)
            self.cell(self.w - 2 * M - 52, 9, f"Ref {ref} · Generated {generated} · Page {self.page_no()} / {{nb}}", new_x="LMARGIN", new_y="NEXT")
            if qr is not None:
                try:
                    self.image(qr, x=self.w - M - 38, y=self.h - 44, w=36, h=36)
                except Exception:  # noqa: BLE001
                    pass

    M = 32.0
    pdf = Doc(orientation="L" if landscape else "P", unit="pt", format="A4")
    pdf.set_margins(M, 40.0, M)
    pdf.set_auto_page_break(True, margin=58.0)
    pdf.set_title(f"{title_en} — {company.get('name', '')}")
    pdf.set_author(company.get("name", ""))
    pdf.alias_nb_pages("{nb}")
    regular, bold = _FONTS / "FreeSans.ttf", _FONTS / "FreeSansBold.ttf"
    pdf.add_font("rpt", "", str(regular))
    pdf.add_font("rpt", "B", str(bold))
    pdf.add_font("rptbn", "", str(regular))  # bold face lacks Bangla/৳ — borrow from regular
    pdf.set_fallback_fonts(["rptbn"], exact_match=False)
    pdf.set_text_shaping(True)
    pdf.add_page()

    # ── Brand band ──
    cw = pdf.w - 2 * M
    band = 74.0
    pdf.set_fill_color(*BRAND)
    pdf.rect(M, 28, cw, band, style="F", round_corners=True, corner_radius=8)
    tx = M + 16
    if logo:
        try:
            pdf.image(str(logo), x=M + 12, y=28 + (band - 46) / 2, w=46, h=46)
            tx = M + 12 + 46 + 12
        except Exception:  # noqa: BLE001
            pass
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("rpt", "B", 16)
    pdf.set_xy(tx, 40)
    pdf.cell(cw * 0.6, 20, company.get("name", ""), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("rpt", "", 8)
    pdf.set_text_color(191, 219, 254)
    pdf.set_xy(tx, 62)
    pdf.cell(cw * 0.6, 11, company.get("tagline", ""), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("rpt", "B", 9)
    pdf.set_xy(M + cw - 200, 40)
    pdf.cell(188, 12, "BUSINESS REPORT", align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("rpt", "", 8)
    pdf.set_xy(M + cw - 200, 56)
    pdf.cell(188, 11, f"Ref {ref}", align="R", new_x="LMARGIN", new_y="NEXT")

    # ── Title + period ──
    pdf.set_xy(M, 114)
    pdf.set_text_color(*BRAND_DARK)
    pdf.set_font("rpt", "B", 17)
    pdf.cell(0, 22, f"{title_en}  ·  {title_bn}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("rpt", "", 9)
    pdf.set_text_color(*MUTED)
    pdf.cell(0, 13, f"Period: {period}   ·   Generated {generated}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # ── Summary strip ──
    boxes = [("Records", f"{len(rows):,}")] + [(_pretty(k), _fmt(k, v)) for k, v in _totals(headers, rows)]
    bw = (cw - 8 * (len(boxes) - 1)) / len(boxes)
    y0 = pdf.get_y()
    for i, (label, val) in enumerate(boxes):
        x = M + i * (bw + 8)
        pdf.set_fill_color(*ROW_ALT)
        pdf.set_draw_color(*GRID)
        pdf.rect(x, y0, bw, 38, style="DF", round_corners=True, corner_radius=6)
        pdf.set_xy(x + 8, y0 + 6)
        pdf.set_font("rpt", "", 7.5)
        pdf.set_text_color(*MUTED)
        pdf.cell(bw - 16, 10, label)
        pdf.set_xy(x + 8, y0 + 18)
        pdf.set_font("rpt", "B", 11)
        pdf.set_text_color(*INK)
        pdf.cell(bw - 16, 14, val)
    pdf.set_y(y0 + 50)

    # ── Table ──
    n = len(headers)
    pdf.set_text_color(*INK)
    pdf.set_font("rpt", "", 8)
    pdf.set_draw_color(*GRID)
    pdf.set_line_width(0.4)
    with pdf.table(
        col_widths=tuple([cw / n] * n),
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=BRAND, size_pt=8.5),
        cell_fill_color=ROW_ALT, cell_fill_mode="ROWS", line_height=11, padding=4,
        text_align=tuple("RIGHT" if any(_num(r.get(h)) is not None for r in rows) else "LEFT" for h in headers),
        repeat_headings=1, first_row_as_headings=True,
    ) as table:
        head = table.row()
        for h in headers:
            head.cell(_pretty(h))
        for r in rows:
            row = table.row()
            for h in headers:
                row.cell(_fmt(h, r.get(h)))
        if not rows:
            row = table.row()
            row.cell("No records in this period", colspan=n)
    return bytes(pdf.output())


def build_report_xlsx(report_type: str, rows: list[dict[str, Any]], period: str, company: dict[str, str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    title_en, title_bn = TITLES.get(report_type, (f"{_pretty(report_type)} Report", "রিপোর্ট"))
    headers = _headers(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = title_en[:28]
    brand = PatternFill("solid", fgColor="1D4ED8")
    alt = PatternFill("solid", fgColor="F3F6FD")
    thin = Side(style="thin", color="DBE3F0")

    top = 1
    logo = _logo()
    if logo:
        try:
            img = XLImage(str(logo)); img.width, img.height = 54, 54
            ws.add_image(img, "A1")
        except Exception:  # noqa: BLE001
            pass
    ws["B1"] = company.get("name", ""); ws["B1"].font = Font(bold=True, size=16, color="1E3A8A")
    ws["B2"] = company.get("tagline", ""); ws["B2"].font = Font(size=9, color="6B7280")
    ws["B3"] = f"{title_en} · {title_bn}"; ws["B3"].font = Font(bold=True, size=13)
    ws["B4"] = f"Period: {period}   ·   Generated {datetime.now(timezone.utc):%d %b %Y, %H:%M UTC}   ·   {len(rows)} records"
    ws["B4"].font = Font(size=9, color="6B7280")
    for r in (1, 2, 3, 4):
        ws.row_dimensions[r].height = 20
    top = 6
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=top, column=c, value=_pretty(h))
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = brand
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(rows, 1):
        for c, h in enumerate(headers, 1):
            v = r.get(h)
            if isinstance(v, Decimal):
                v = float(v)
            if isinstance(v, (datetime, date)):
                v = v.isoformat()
            cell = ws.cell(row=top + i, column=c, value=v)
            cell.border = Border(bottom=thin)
            if i % 2 == 0:
                cell.fill = alt
            if isinstance(v, float):
                cell.number_format = '"৳"#,##0.00' if _MONEY.search(h) else "#,##0.00"
            elif isinstance(v, int) and not isinstance(v, bool):
                cell.number_format = "#,##0"
    # totals row for money / count columns
    tot = top + len(rows) + 1
    ws.cell(row=tot, column=1, value="Total").font = Font(bold=True)
    for c, h in enumerate(headers, 1):
        if any(h == k for k, _ in _totals(headers, rows)):
            col = get_column_letter(c)
            cell = ws.cell(row=tot, column=c, value=f"=SUM({col}{top + 1}:{col}{top + len(rows)})")
            cell.font = Font(bold=True)
            cell.number_format = '"৳"#,##0.00' if _MONEY.search(h) else "#,##0"
    for c, h in enumerate(headers, 1):
        width = max([len(_pretty(h))] + [len(str(r.get(h, ""))) for r in rows[:200]]) + 3
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 12), 44)
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    ws.sheet_view.showGridLines = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
